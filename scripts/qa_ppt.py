"""PPTX 程序化 QA：文本溢出估算 / 越界 / 文本框重叠检查"""
import sys
from pptx import Presentation
from pptx.util import Emu

PATH = "KontainKeeper-项目架构设计方案.pptx"
EMU_IN = 914400.0


def est_text_w(text, fs):
    w = 0.0
    for ch in text:
        w += (fs / 72.0) * (1.02 if ord(ch) > 0x2E7F else 0.55)
    return w


def est_lines(text, box_w, fs):
    lines = 0
    for para in text.split("\n"):
        if not para.strip():
            lines += 1
            continue
        lines += max(1, int(est_text_w(para, fs) / max(box_w, 0.1)) + (1 if est_text_w(para, fs) % max(box_w, 0.1) > 0.01 else 0))
    return lines


def walk(shapes, out):
    for sh in shapes:
        if sh.shape_type == 6:
            walk(sh.shapes, out)
        elif sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh)


prs = Presentation(PATH)
sw, sh_ = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
issues = []
for idx, slide in enumerate(prs.slides, 1):
    boxes = []
    walk(slide.shapes, boxes)
    for b in boxes:
        x, y = b.left / EMU_IN, b.top / EMU_IN
        w, h = b.width / EMU_IN, b.height / EMU_IN
        txt = b.text_frame.text
        fs = 18.0
        for para in b.text_frame.paragraphs:
            for run in para.runs:
                if run.font.size:
                    fs = run.font.size.pt
                    break
            break
        # 越界
        if x < -0.01 or y < -0.01 or x + w > sw + 0.01 or y + h > sh_ + 0.01:
            issues.append(f"S{idx} 越界: ({x:.2f},{y:.2f},{w:.2f},{h:.2f}) {txt[:18]!r}")
        # 溢出估算（宽度按单行最长段、高度按行数）
        longest = max((p for p in txt.split("\n")), key=lambda s: est_text_w(s, fs), default="")
        est_w = est_text_w(longest, fs)
        if est_w > w * 1.06 and est_w <= w * 3.0:  # 明显超宽（>6%）
            issues.append(f"S{idx} 可能超宽: est {est_w:.2f}in > box {w:.2f}in fs={fs:.0f} {txt[:22]!r}")
        n_lines = est_lines(txt, w, fs)
        est_h = n_lines * fs / 72.0 * 1.35
        if est_h > h * 1.18 and h > 0.15:
            issues.append(f"S{idx} 可能超高: est {est_h:.2f}in > box {h:.2f}in lines={n_lines} fs={fs:.0f} {txt[:22]!r}")
    # 重叠（文本框之间）
    rects = [(b, b.left / EMU_IN, b.top / EMU_IN, b.width / EMU_IN, b.height / EMU_IN) for b in boxes]
    for i in range(len(rects)):
        for j in range(i + 1, len(rects)):
            b1, x1, y1, w1, h1 = rects[i]
            b2, x2, y2, w2, h2 = rects[j]
            ox = min(x1 + w1, x2 + w2) - max(x1, x2)
            oy = min(y1 + h1, y2 + h2) - max(y1, y2)
            if ox > 0.08 and oy > 0.08:
                a1, a2 = w1 * h1, w2 * h2
                # 忽略小框在大框内的合理嵌套（如芯片在 band 内）
                if min(a1, a2) / max(a1, a2) < 0.4 and (ox * oy) / min(a1, a2) > 0.85:
                    continue
                issues.append(f"S{idx} 重叠: {b1.text_frame.text[:12]!r} × {b2.text_frame.text[:12]!r} ov=({ox:.2f},{oy:.2f})")

print(f"slides={len(prs.slides)} issues={len(issues)}")
for it in issues:
    print(" -", it)
